import os
import sys

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT_DIR)

import streamlit as st
import pandas as pd
from io import BytesIO
from datetime import datetime
import json
import requests

from openpyxl import Workbook
from docx import Document
from pypdf import PdfReader

# ===============================
# CONFIG API URL
# ===============================
API_URL = "http://localhost:8000/generate"
# Khi dùng ngrok → đổi thành:
# API_URL = "https://abc123.ngrok-free.app/generate"

# ===============================
# STATE
# ===============================
if "history" not in st.session_state:
    st.session_state.history = []

if "generated" not in st.session_state:
    st.session_state.generated = False

if "last_result" not in st.session_state:
    st.session_state.last_result = None

# ===============================
# UI
# ===============================
st.set_page_config(page_title="AI Test Case Generator")
st.title("🚀 AI Test Case Generator")

# ===============================
# INPUT
# ===============================
input_type = st.selectbox(
    "Loại yêu cầu", ["User Story", "Use Case Spec", "Natural Language"]
)

language = st.selectbox("Ngôn ngữ", ["Tiếng Việt", "English"])

option = st.radio("Input", ["Text", "Upload File"])


def get_placeholder(t):
    if t == "User Story":
        return "As a user, I want to login..."
    elif t == "Use Case Spec":
        return "Use Case: Login..."
    return "User logs into system..."


def read_file(file):
    ext = file.name.split(".")[-1]

    if ext == "txt":
        return file.read().decode("utf-8")
    elif ext == "docx":
        return "\n".join([p.text for p in Document(file).paragraphs])
    elif ext == "pdf":
        reader = PdfReader(file)
        return "\n".join([p.extract_text() or "" for p in reader.pages])
    return ""


user_input = ""

if option == "Text":
    user_input = st.text_area(
        "Requirement", height=250, placeholder=get_placeholder(input_type)
    )
else:
    f = st.file_uploader("Upload", type=["txt", "docx", "pdf"])
    if f:
        user_input = read_file(f)
        st.text_area("Preview", user_input, height=200)


# ===============================
# EXPORT FUNCTIONS
# ===============================
def export_excel(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases"

    headers = [
        "id",
        "feature",
        "screen",
        "description",
        "pre_condition",
        "steps",
        "expected_result",
    ]

    ws.append(headers)

    for row in data:
        ws.append(
            [
                row.get("id"),
                row.get("feature"),
                row.get("screen"),
                row.get("description"),
                row.get("pre_condition"),
                "\n".join(row.get("steps", [])),
                row.get("expected_result"),
            ]
        )

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def export_testdata_excel(testdata):
    df = pd.DataFrame(testdata)
    output = BytesIO()
    df.to_excel(output, index=False)
    return output.getvalue()


def generate_filename():
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    return f"TestCases_{timestamp}.xlsx"


# ===============================
# GENERATE BUTTON (CALL API)
# ===============================
if st.button("Generate"):

    if not user_input.strip():
        st.warning("Nhập nội dung")
        st.stop()

    with st.spinner("Generating..."):
        try:
            response = requests.post(
                API_URL,
                json={
                    "text": user_input,
                    "input_type": input_type,
                    "language": language,
                },
                timeout=60,
            )

            data = response.json()

        except Exception as e:
            st.error(f"API error: {e}")
            st.stop()

    if data.get("status") != "success":
        st.error("Generation failed")
    else:
        st.session_state.generated = True
        st.session_state.last_result = data
        st.session_state.history.append(data)

# ===============================
# SHOW RESULT
# ===============================
if st.session_state.generated and st.session_state.last_result:

    result = st.session_state.last_result

    tab1, tab2, tab3 = st.tabs(["Analysis", "TestCases", "Test Data"])

    # Analysis
    with tab1:
        st.json(result.get("analysis"))

    # TestCases
    with tab2:
        df = pd.DataFrame(result.get("test_cases", []))
        st.dataframe(df)

        excel = export_excel(result.get("test_cases", []))
        st.download_button("📥 Download TestCases", excel, generate_filename())

    # Test Data
    with tab3:
        testdata = result.get("test_data", [])

        if testdata:
            td_df = pd.DataFrame(testdata)
            st.dataframe(td_df)

            json_bytes = json.dumps(testdata, indent=2).encode("utf-8")

            st.download_button("📥 Download JSON", json_bytes, "test_data.json")

            excel_td = export_testdata_excel(testdata)
            st.download_button("📥 Download Excel", excel_td, "test_data.xlsx")
        else:
            st.warning("No test data generated")

# ===============================
# HISTORY
# ===============================
if st.session_state.history:
    st.subheader("History")

    for i, item in enumerate(reversed(st.session_state.history)):
        with st.expander(f"Result {i+1}"):

            st.text(item.get("input", ""))

            df = pd.DataFrame(item.get("test_cases", []))
            st.dataframe(df)

            excel = export_excel(item.get("test_cases", []))
            st.download_button(
                "📥 Download TestCases",
                excel,
                f"testcases_{i}.xlsx",
                key=f"his_tc_{i}",
            )
